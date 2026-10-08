#!/usr/bin/env python3
"""
Kleinanzeigen Häuser-Scraper -> Excel

Sammelt Haus-Inserate (Kauf + Miete, deutschlandweit, private + gewerbliche
Anbieter) von kleinanzeigen.de und speichert NUR Inserate, die
  - Straße UND Hausnummer,
  - Baujahr und
  - Gebäudetyp (Haustyp)
im Inserat selbst angeben. Es werden keine Daten ergänzt oder geschätzt:
Jeder Wert stammt 1:1 aus dem Inserat, die Quelle (Link) steht in jeder Zeile.

Benutzung (Details siehe README.md):
    pip install -r requirements.txt
    python kleinanzeigen_haeuser.py --test       # Probelauf: 10 Treffer
    python kleinanzeigen_haeuser.py              # Standard: 500 Treffer
    python kleinanzeigen_haeuser.py --ziel 100   # weniger Treffer
    python kleinanzeigen_haeuser.py --nur-kauf   # nur Häuser zum Kauf

Das Skript kann jederzeit mit Strg+C abgebrochen und später erneut gestartet
werden – bereits geprüfte Inserate werden in fortschritt.json gemerkt.
"""

import argparse
import json
import random
import re
import sys
import time
import warnings
from pathlib import Path

# Harmloser Hinweis von urllib3 bei älteren Python-Versionen auf macOS (LibreSSL)
warnings.filterwarnings("ignore", message=".*OpenSSL.*")

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BASIS = "https://www.kleinanzeigen.de"

# Kategorien: Häuser zum Kauf (c208), Häuser zur Miete (c205).
# Die Kategorienummer steckt auch in jedem Inserat-Link (.../1234567890-208-1548).
KATEGORIEN = {
    "kauf": ("s-haus-kaufen", "c208", "208"),
    "miete": ("s-haus-mieten", "c205", "205"),
}

# Kleinanzeigen zeigt pro Suche max. ~50 Seiten. Damit deutschlandweit genug
# Inserate erreichbar sind, wird die Suche in Preisbänder aufgeteilt.
PREISBAENDER_KAUF = [
    (0, 100000), (100001, 150000), (150001, 200000), (200001, 250000),
    (250001, 300000), (300001, 350000), (350001, 400000), (400001, 500000),
    (500001, 650000), (650001, 850000), (850001, 1200000), (1200001, None),
]
PREISBAENDER_MIETE = [
    (0, 800), (801, 1100), (1101, 1400), (1401, 1800), (1801, 2500), (2501, None),
]

MAX_SEITEN = 50
PAUSE_MIN, PAUSE_MAX = 2.5, 5.0  # Sekunden zwischen Anfragen (fair bleiben)

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Accept-Language": "de-DE,de;q=0.9",
}

SPALTEN = ["Stadt", "PLZ", "Straße", "Hausnummer", "Gebäudetyp", "Baujahr",
           "Wohnfläche (m²)", "Sonstige Informationen", "Quelle"]

# Detail-Felder, die in "Sonstige Informationen" landen (falls vorhanden)
SONSTIGES_FELDER = [
    "Heizungsart", "Wesentlicher Energieträger", "Energieträger",
    "Energieeffizienzklasse", "Endenergiebedarf", "Endenergieverbrauch",
    "Energieverbrauch", "Energieausweis", "Energieausweistyp",
    "Grundstücksfläche", "Zimmer", "Badezimmer", "Etagen", "Schlafzimmer",
    "Bezugsfrei ab", "Provision", "Denkmalschutzobjekt",
]


# Alle Dateien (Excel, Fortschritt) landen im Ordner dieses Skripts
ORDNER = Path(__file__).resolve().parent


class Gesperrt(Exception):
    pass


class ZielErreicht(Exception):
    pass


def warte():
    time.sleep(random.uniform(PAUSE_MIN, PAUSE_MAX))


def lade(session, url):
    """Lädt eine Seite. Bei Sperre/Captcha wird abgebrochen statt umgangen."""
    for versuch in range(3):
        try:
            r = session.get(url, headers=HEADERS, timeout=30)
        except requests.RequestException as e:
            print(f"  Netzwerkfehler ({e}), neuer Versuch in 30 s ...")
            time.sleep(30)
            continue
        if r.status_code == 404:
            return None
        if r.status_code in (403, 429):
            if versuch < 2:
                print(f"  HTTP {r.status_code} – Pause 5 Minuten ...")
                time.sleep(300)
                continue
            raise Gesperrt(f"HTTP {r.status_code} bei {url}")
        r.raise_for_status()
        if ("captcha" in r.text.lower() and "/s-anzeige/" not in r.text
                and "viewad-title" not in r.text):
            raise Gesperrt("Captcha-Seite erhalten")
        return r.text
    raise Gesperrt(f"Seite nicht ladbar: {url}")


def such_url(slug, cat, seite, preis):
    teile = [slug]
    if preis:
        lo, hi = preis
        teile.append(f"preis:{lo}:{hi if hi is not None else ''}")
    if seite > 1:
        teile.append(f"seite:{seite}")
    teile.append(cat)
    return f"{BASIS}/" + "/".join(teile)


# Inserat-Links, z.B. /s-anzeige/titel-des-inserats/3528736838-208-1548
ANZEIGE_RE = re.compile(r"/s-anzeige/[^\"'\s?#<>]+/\d+-(\d+)-\d+")


def inserat_links(html, katnr):
    """Alle Inserat-Links der Suchseite aus der gewünschten Kategorie (ohne Duplikate)."""
    links = []
    for m in ANZEIGE_RE.finditer(html):
        url = BASIS + m.group(0)
        if m.group(1) == katnr and url not in links:
            links.append(url)
    return links


def text(el):
    return " ".join(el.get_text(" ", strip=True).split()) if el else ""


STRASSE_RE = re.compile(
    r"^(?P<strasse>.*?[A-Za-zÄÖÜäöüß.\-]\s*)\s+(?P<nr>\d+\s*[a-zA-Z]?(\s*[-/]\s*\d+\s*[a-zA-Z]?)?)$"
)


def zerlege_strasse(roh):
    roh = roh.strip().rstrip(",").strip()
    m = STRASSE_RE.match(roh)
    if not m:
        return None, None
    return m.group("strasse").strip(), m.group("nr").replace(" ", "")


def parse_inserat(html, url):
    soup = BeautifulSoup(html, "html.parser")

    # Adresse
    strasse_roh = text(soup.select_one("#street-address")
                       or soup.select_one("[itemprop='streetAddress']"))
    if not strasse_roh:
        m = re.search(r'"streetAddress"\s*:\s*"([^"]+)"', html)
        strasse_roh = m.group(1) if m else ""
    strasse, nr = zerlege_strasse(strasse_roh) if strasse_roh else (None, None)

    ort_roh = text(soup.select_one("#viewad-locality"))  # z.B. "12345 Berlin - Mitte"
    plz, stadt = None, None
    m = re.match(r"(\d{5})\s+(.*)", ort_roh)
    if m:
        plz, stadt = m.group(1), m.group(2).strip()

    # Detail-Liste (Schlüssel -> Wert)
    details = {}
    for li in soup.select("li.addetailslist--detail"):
        wert_el = li.select_one(".addetailslist--detail--value")
        wert = text(wert_el)
        if wert_el:
            wert_el.extract()
        schluessel = text(li).rstrip(":")
        if schluessel:
            details[schluessel] = wert

    # Ausstattungs-Tags (z.B. "Keller", "Garten")
    tags = [text(t) for t in soup.select(".checktag")]

    preis = text(soup.select_one("#viewad-price"))

    wohnflaeche = details.get("Wohnfläche", "")
    wohnflaeche = re.sub(r"\s*m²\s*$", "", wohnflaeche).strip()

    sonstiges = []
    if preis:
        sonstiges.append(f"Preis: {preis}")
    for feld in SONSTIGES_FELDER:
        if details.get(feld):
            sonstiges.append(f"{feld}: {details[feld]}")
    if tags:
        sonstiges.append("Ausstattung: " + ", ".join(tags))

    return {
        "Stadt": stadt,
        "PLZ": plz,
        "Straße": strasse,
        "Hausnummer": nr,
        "Gebäudetyp": details.get("Haustyp") or details.get("Typ"),
        "Baujahr": details.get("Baujahr"),
        "Wohnfläche (m²)": wohnflaeche or None,
        "Sonstige Informationen": " | ".join(sonstiges),
        "Quelle": url,
    }


def vollstaendig(z):
    return all(z.get(k) for k in ("Straße", "Hausnummer", "Baujahr", "Gebäudetyp", "PLZ", "Stadt"))


def speichere_excel(zeilen, pfad):
    wb = Workbook()
    ws = wb.active
    ws.title = "Häuser"
    ws.append(SPALTEN)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    for z in zeilen:
        ws.append([z.get(s) for s in SPALTEN])
        zelle = ws.cell(row=ws.max_row, column=len(SPALTEN))
        zelle.hyperlink = z["Quelle"]
        zelle.font = Font(color="0563C1", underline="single")
    breiten = [18, 8, 28, 11, 20, 9, 15, 80, 60]
    for i, b in enumerate(breiten, 1):
        ws.column_dimensions[get_column_letter(i)].width = b
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(pfad)


def main():
    ap = argparse.ArgumentParser(description="Kleinanzeigen Häuser -> Excel")
    ap.add_argument("--ziel", type=int, default=500, help="Anzahl vollständiger Treffer (Standard 500)")
    ap.add_argument("--nur-kauf", action="store_true", help="nur Häuser zum Kauf")
    ap.add_argument("--nur-miete", action="store_true", help="nur Häuser zur Miete")
    ap.add_argument("--datei", default="haeuser_kleinanzeigen.xlsx", help="Excel-Ausgabedatei")
    ap.add_argument("--test", action="store_true",
                    help="Probelauf mit 10 Treffern in eigene Dateien (test_*.xlsx / fortschritt_test.json)")
    args = ap.parse_args()
    if args.test:
        args.ziel = min(args.ziel, 10)
        args.datei = "test_" + args.datei

    # Windows-Konsolen können sonst an Sonderzeichen wie ✔ scheitern
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace", line_buffering=True)
    except (AttributeError, ValueError):
        pass

    excel_pfad = ORDNER / args.datei
    fortschritt_pfad = ORDNER / ("fortschritt_test.json" if args.test else "fortschritt.json")
    stand = {"geprueft": [], "treffer": []}
    if fortschritt_pfad.exists():
        stand = json.loads(fortschritt_pfad.read_text(encoding="utf-8"))
        print(f"Fortsetzung: {len(stand['treffer'])} Treffer, "
              f"{len(stand['geprueft'])} Inserate bereits geprüft.")
    geprueft = set(stand["geprueft"])
    treffer = stand["treffer"]

    def sichern():
        stand["geprueft"] = sorted(geprueft)
        stand["treffer"] = treffer
        fortschritt_pfad.write_text(json.dumps(stand, ensure_ascii=False), encoding="utf-8")
        try:
            speichere_excel(treffer, excel_pfad)
        except PermissionError:
            print(f"  Hinweis: {excel_pfad.name} ist gerade geöffnet und kann nicht gespeichert "
                  "werden – bitte Excel schließen. Der Fortschritt ist trotzdem gesichert.")

    arten = ["kauf", "miete"]
    if args.nur_kauf:
        arten = ["kauf"]
    elif args.nur_miete:
        arten = ["miete"]

    session = requests.Session()
    try:
        for art in arten:
            slug, cat, katnr = KATEGORIEN[art]
            baender = PREISBAENDER_KAUF if art == "kauf" else PREISBAENDER_MIETE
            for band in baender:
                vorherige = None
                for seite in range(1, MAX_SEITEN + 1):
                    if len(treffer) >= args.ziel:
                        raise ZielErreicht
                    url = such_url(slug, cat, seite, band)
                    print(f"[{art} {band}] Seite {seite}: {url}")
                    html = lade(session, url)
                    warte()
                    links = inserat_links(html, katnr) if html else []
                    if not links or links == vorherige:
                        break  # keine weiteren Seiten in diesem Band
                    vorherige = links
                    neue = [l for l in links if l not in geprueft]
                    vorher = len(treffer)
                    for link in neue:
                        if len(treffer) >= args.ziel:
                            raise ZielErreicht
                        detail = lade(session, link)
                        warte()
                        geprueft.add(link)
                        if not detail:
                            continue
                        z = parse_inserat(detail, link)
                        if vollstaendig(z):
                            treffer.append(z)
                            print(f"   ✔ {len(treffer)}/{args.ziel}: {z['Straße']} {z['Hausnummer']}, "
                                  f"{z['PLZ']} {z['Stadt']} ({z['Gebäudetyp']}, {z['Baujahr']})")
                    print(f"   {len(links)} Inserate auf der Seite, {len(neue)} neu geprüft, "
                          f"{len(treffer) - vorher} vollständig (gesamt {len(treffer)}/{args.ziel})")
                    sichern()
    except ZielErreicht:
        pass
    except Gesperrt as e:
        print(f"\nKleinanzeigen ist nicht erreichbar oder blockiert gerade die Anfragen ({e}).")
        print("Bitte später (z.B. in ein paar Stunden) erneut starten – der Fortschritt ist gespeichert.")
    except KeyboardInterrupt:
        print("\nAbgebrochen – Fortschritt wird gespeichert.")
    finally:
        sichern()
        print(f"\n{len(treffer)} vollständige Häuser gespeichert in: {excel_pfad}")


if __name__ == "__main__":
    sys.exit(main())
